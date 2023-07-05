import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows 
from openpyxl.chart import BarChart, Series, Reference

file_name = '/workspaces/1202/video_game_sales.csv'

df = pd.read_csv(file_name)
#print(df.head())

#how many game per platform, create a function

def game_per_platform(_df):
    games = _df.groupby(['Platform'])['Name'].count().reset_index()
    return games
print(game_per_platform(df))

games_df = game_per_platform(df)
games_df = games_df.sort_values('Name',ascending=False)
games_df = games_df.['Platform', 'Number_of_games']

wb = Workbook()
ws = wb.active

for r in dataframe_to_rows(games_df, index=False, header=True):
    ws.append(r)

last_row = len(ws['A'])
chart1 = BarChart()
chart1.type = "col"
chart1.style = 10
chart1.title = "Vg_sales"
chart1.y_axis.title = 'Count'
chart1.x_axis.title = 'Platform'

data = Reference(ws, min_col=2, min_row=1, max_row=last_row, max_col=2)
cats = Reference(ws, min_col=1, min_row=2, max_row=last_row)
chart1.add_data(data, titles_from_data=True)
chart1.set_categories(cats)
chart1.shape = 4
ws.add_chart(chart1, "E3")

wb.save('week9/week9.xlsx')